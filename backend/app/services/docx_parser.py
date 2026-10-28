"""DOCX text extraction with python-docx."""
import io
import logging

from docx import Document

from app.core.exceptions import CorruptedDocumentError, EmptyDocumentError, format_safe_traceback

logger = logging.getLogger(__name__)

# Rough estimate: python-docx cannot know the rendered page count.
_WORDS_PER_PAGE = 500


def extract_docx(data: bytes) -> dict:
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        logger.warning("DOCX could not be opened:\n%s", format_safe_traceback(exc))
        raise CorruptedDocumentError("The DOCX file could not be read. It may be corrupted.") from exc

    parts: list[str] = [p.text for p in doc.paragraphs]
    # Many resume templates put content inside tables (two-column layouts).
    for table in doc.tables:
        for row in table.rows:
            seen = set()
            for cell in row.cells:
                if id(cell._tc) in seen:  # merged cells repeat
                    continue
                seen.add(id(cell._tc))
                parts.extend(p.text for p in cell.paragraphs)

    text = "\n".join(parts).strip()
    if not text:
        raise EmptyDocumentError()
    words = len(text.split())
    return {"text": text, "page_count": max(1, round(words / _WORDS_PER_PAGE)), "file_type": "docx"}
